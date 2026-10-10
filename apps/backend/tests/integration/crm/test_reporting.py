"""Home page (REP-01), firm dashboard (REP-02), fee forecast (REP-03), win/loss (REP-04),
reports and exports behind the wall (AUTHZ-02 c4), Read-only sees aggregates only (ROLE-01 c3)."""

import io
import time
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from src.infrastructure.clock import FrozenClock, SystemClock, set_clock
from src.modules.role.permissions import RoleName
from tests.integration.crm.deal_helpers import build_world, create_mandate
from tests.integration.foundation.helpers import step_up

pytestmark = pytest.mark.asyncio

NOW = datetime(2026, 10, 14, 9, 0, tzinfo=UTC)  # a Wednesday
REPORTS = "/api/v1/reports"


@pytest.fixture
async def world(client, seeded_db):
    set_clock(FrozenClock(NOW))
    built = await build_world(seeded_db)
    built.partner = await built.actor("paula", RoleName.PARTNER)
    built.analyst = await built.actor("anna", RoleName.ANALYST)
    built.outsider = await built.actor("otto", RoleName.PARTNER)
    built.backoffice = await built.actor("cora", RoleName.COMPLIANCE)
    built.reader = await built.actor("rita", RoleName.READ_ONLY)
    built.client_id = await built.company("Acme Holding")
    yield built
    await built.close()
    set_clock(SystemClock())


async def get(actor, path: str, **params):
    return await actor.client.get(path, params=params, headers=actor.headers)


async def export_(actor, path: str):
    return await actor.client.get(path, headers=step_up(actor.email, actor.headers))


async def put(actor, path: str, body: dict):
    return await actor.client.put(path, json=body, headers=actor.headers)


async def post(actor, path: str, body: dict):
    return await actor.client.post(path, json=body, headers=actor.headers)


async def deal(world, code: str, lead=None, mandate_type: str = "sell_side", ev: str = "25000") -> int:
    lead = lead or world.partner
    # Each deal its own client: a shared client would be a potential conflict and hide the deal (AUTHZ-05 c1).
    client_id = await world.company(code[::-1].capitalize())
    mandate = await create_mandate(world, lead, client_id, code_name=code, type=mandate_type, estimated_ev=ev)
    return mandate["id"]


async def fees(world, mandate_id: int, lead=None, retainer: str = "10000", percent: str = "2") -> None:
    response = await put(
        lead or world.partner,
        f"/api/v1/mandates/{mandate_id}/fees",
        {"retainer_monthly": retainer, "success_fee_kind": "percent", "success_fee_percent": percent},
    )
    assert response.status_code == 200, response.text


async def forecast(world, mandate_id: int, probability: int | None, closing: str | None, lead=None) -> dict:
    response = await put(
        lead or world.partner,
        f"/api/v1/mandates/{mandate_id}/forecast",
        {"probability": probability, "expected_closing_date": closing},
    )
    assert response.status_code == 200, response.text
    return response.json()


async def staff(world, mandate_id: int, actor, role: str = "analyst", lead=None) -> None:
    lead = lead or world.partner
    response = await post(lead, f"/api/v1/mandates/{mandate_id}/team", {"user_id": actor.id, "role": role})
    assert response.status_code == 201, response.text


# REP-01


async def test_home_shows_my_active_deals_by_stage_tasks_due_this_week_and_key_dates(world):
    falcon = await deal(world, "Falcon")
    await deal(world, "Heron", lead=world.outsider)
    await staff(world, falcon, world.analyst)
    await post(world.partner, f"/api/v1/mandates/{falcon}/milestones", {"name": "Teaser out", "planned_date": "2026-10-20"})
    await post(world.partner, f"/api/v1/mandates/{falcon}/milestones", {"name": "Far away", "planned_date": "2027-03-01"})
    await forecast(world, falcon, None, "2026-11-01")
    for title, due in (("Call the client", "2026-10-16"), ("Next week", "2026-10-25")):
        response = await post(
            world.partner, "/api/v1/tasks/", {"title": title, "assignee_id": world.analyst.id, "due_date": due}
        )
        assert response.status_code == 201, response.text

    started = time.perf_counter()
    response = await get(world.analyst, f"{REPORTS}/home")
    elapsed = time.perf_counter() - started

    assert response.status_code == 200, response.text
    assert elapsed < 2
    home = response.json()
    assert home["is_new"] is False
    stages = {group["stage"]: [d["code_name"] for d in group["deals"]] for group in home["deals_by_stage"]}
    assert sum(stages.values(), []) == ["Falcon"]
    assert [task["title"] for task in home["tasks_due_this_week"]] == ["Call the client"]
    assert (home["week_start"], home["week_end"]) == ("2026-10-12", "2026-10-18")
    assert [(k["code_name"], k["name"]) for k in home["key_dates"]] == [
        ("Falcon", "Teaser out"),
        ("Falcon", "Expected closing"),
    ]
    assert "Heron" not in response.text


async def test_home_of_a_new_user_flags_the_empty_state(world):
    response = await get(world.analyst, f"{REPORTS}/home")

    assert response.status_code == 200
    assert response.json() == {
        "deals_by_stage": [],
        "tasks_due_this_week": [],
        "key_dates": [],
        "week_start": "2026-10-12",
        "week_end": "2026-10-18",
        "is_new": True,
    }


# REP-02


async def test_dashboard_counts_mandates_by_stage_and_type_and_weights_fees_by_quarter(world):
    falcon = await deal(world, "Falcon")
    heron = await deal(world, "Heron", mandate_type="buy_side", ev="50000")
    await fees(world, falcon)  # 2% of 25m = 500k
    await fees(world, heron)  # 2% of 50m = 1m
    await forecast(world, falcon, 50, "2026-11-15")
    await forecast(world, heron, 10, "2027-02-01")

    response = await get(world.partner, f"{REPORTS}/dashboard")

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total_mandates"] == 2
    cells = {(c["stage"], c["type"]): c["count"] for c in body["by_stage_and_type"]}
    assert sum(cells.values()) == 2
    assert {kind for _, kind in cells} == {"sell_side", "buy_side"}
    quarters = {q["quarter"]: q for q in body["weighted_fees_by_quarter"]}
    assert Decimal(quarters["2026-Q4"]["weighted_fees"]) == Decimal("250000.00")
    assert Decimal(quarters["2027-Q1"]["weighted_fees"]) == Decimal("100000.00")
    assert Decimal(quarters["2027-Q1"]["expected_fees"]) == Decimal("1000000.00")


async def test_forecast_probability_defaults_to_the_stage(world):
    falcon = await deal(world, "Falcon")
    read = (await get(world.partner, f"/api/v1/mandates/{falcon}/forecast")).json()
    assert read["probability_is_default"] is True
    assert 0 < read["probability"] < 100
    assert (await put(world.outsider, f"/api/v1/mandates/{falcon}/forecast", {"probability": 90})).status_code == 404


async def test_unstaffed_deals_count_only_in_aggregates_never_by_name(world):
    lone = await deal(world, "Kestrel", lead=world.outsider)
    await fees(world, lone, lead=world.outsider)
    await forecast(world, lone, 100, "2026-11-15", lead=world.outsider)

    alone = (await get(world.partner, f"{REPORTS}/dashboard")).json()
    # One hidden deal would be identifiable by the total moving: it is left out entirely.
    assert alone == {
        "total_mandates": 0,
        "by_stage_and_type": [],
        "weighted_fees_by_quarter": [],
        "other_teams_weighted_fees": None,
    }

    for code in ("Osprey", "Plover"):
        mandate_id = await deal(world, code, lead=world.outsider)
        await fees(world, mandate_id, lead=world.outsider)
        await forecast(world, mandate_id, 100, "2026-11-20", lead=world.outsider)

    response = await get(world.partner, f"{REPORTS}/dashboard")
    body = response.json()
    # REP-02 c1: a cell or quarter counts hidden deals once it holds HIDDEN_CELL_MINIMUM of them.
    assert body["total_mandates"] == 3
    assert [(c["type"], c["count"]) for c in body["by_stage_and_type"]] == [("sell_side", 3)]
    quarters = {q["quarter"]: q for q in body["weighted_fees_by_quarter"]}
    assert list(quarters) == ["2026-Q4"]
    assert quarters["2026-Q4"]["deals"] == 3
    assert Decimal(quarters["2026-Q4"]["weighted_fees"]) == Decimal("1500000.00")
    assert Decimal(body["other_teams_weighted_fees"]) == Decimal("1500000.00")

    mine = await deal(world, "Falcon", mandate_type="buy_side")
    await fees(world, mine)
    await forecast(world, mine, 100, "2027-02-01")
    mixed = (await get(world.partner, f"{REPORTS}/dashboard")).json()
    assert sorted((c["type"], c["count"]) for c in mixed["by_stage_and_type"]) == [("buy_side", 1), ("sell_side", 3)]
    assert {q["quarter"]: q["deals"] for q in mixed["weighted_fees_by_quarter"]} == {"2026-Q4": 3, "2027-Q1": 1}
    for leak in ("Kestrel", "Osprey", "Plover", "Acme"):
        assert leak not in response.text
    assert "mandate_id" not in response.text
    assert '"id"' not in response.text


# REP-03


async def test_fee_forecast_shows_retainers_and_success_fees_per_month_for_12_months_as_excel(world):
    falcon = await deal(world, "Falcon")
    await fees(world, falcon)  # retainer 10k / month, 500k success fee
    await forecast(world, falcon, 100, "2027-01-15")

    response = await get(world.backoffice, f"{REPORTS}/fee-forecast")

    assert response.status_code == 200, response.text
    months = response.json()["months"]
    assert [m["month"] for m in months][:4] == ["2026-10", "2026-11", "2026-12", "2027-01"]
    assert len(months) == 12
    by_month = {m["month"]: m for m in months}
    assert [Decimal(by_month[k]["retainers"]) for k in ("2026-10", "2027-01", "2027-02")] == [
        Decimal("10000.00"),
        Decimal("10000.00"),
        Decimal("0.00"),
    ]
    assert Decimal(by_month["2027-01"]["success_fees"]) == Decimal("500000.00")
    assert Decimal(by_month["2027-01"]["forecast_total"]) == Decimal("510000.00")

    export = await export_(world.backoffice, f"{REPORTS}/fee-forecast/export")
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    sheet = load_workbook(io.BytesIO(export.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Month"
    assert len(rows) == 13
    assert rows[4][0] == "2027-01"
    assert rows[4][2] == 500000


async def test_fee_forecast_shows_actual_vs_forecast_per_month_from_invoices(world):
    falcon = await deal(world, "Falcon")
    await fees(world, falcon)
    for body in (
        {"type": "retainer", "amount": "10000", "sent_date": "2026-10-01", "paid_date": "2026-11-03"},
        {"type": "retainer", "amount": "10000", "sent_date": "2026-11-01"},
    ):
        response = await post(world.backoffice, f"/api/v1/mandates/{falcon}/invoices", body)
        assert response.status_code == 201, response.text

    months = {m["month"]: m for m in (await get(world.backoffice, f"{REPORTS}/fee-forecast")).json()["months"]}

    assert Decimal(months["2026-10"]["forecast_total"]) == Decimal("10000.00")
    assert Decimal(months["2026-10"]["actual_invoiced"]) == Decimal("10000.00")
    assert Decimal(months["2026-10"]["actual_paid"]) == Decimal("0.00")
    assert Decimal(months["2026-11"]["actual_invoiced"]) == Decimal("10000.00")
    assert Decimal(months["2026-11"]["actual_paid"]) == Decimal("10000.00")


async def test_fee_forecast_needs_fee_visibility(world):
    assert (await get(world.analyst, f"{REPORTS}/fee-forecast")).status_code == 403
    assert (await get(world.analyst, f"{REPORTS}/fee-forecast/export")).status_code == 403


# REP-04


async def test_win_loss_report_for_a_period(world):
    base = "/api/v1/origination/opportunities"

    async def opportunity() -> int:
        body = {
            "company_id": world.client_id,
            "type": "sell_side",
            "trigger": "succession",
            "estimated_ev": "40000",
            "probability": 30,
        }
        response = await post(world.partner, base, body)
        assert response.status_code == 201, response.text
        return response.json()["id"]

    for reason in ("fee", "fee", "timing", "competitor"):
        lost = await post(world.partner, f"{base}/{await opportunity()}/stage", {"stage": "lost", "lost_reason": reason})
        assert lost.status_code == 200, lost.text
    won = await post(world.partner, f"{base}/{await opportunity()}/win", {"code_name": "Swift", "sector_id": world.sector_id})
    assert won.status_code == 201, won.text

    set_clock(FrozenClock(datetime(2026, 7, 16, 9, 0, tzinfo=UTC)))
    closed = await deal(world, "Falcon")
    set_clock(FrozenClock(NOW))
    current = (await get(world.partner, f"/api/v1/mandates/{closed}")).json()
    response = await post(
        world.partner,
        f"/api/v1/mandates/{closed}/close",
        {
            "version": current["version"],
            "outcome": "closed_won",
            "final_ev": "30000",
            "buyer": "Northwind",
            "ev_ebitda": "8.5",
            "closing_date": "2026-10-14",
            "tombstone_allowed": False,
        },
    )
    assert response.status_code == 200, response.text

    report = await get(world.partner, f"{REPORTS}/win-loss", **{"from": "2026-10-01", "to": "2026-10-31"})

    assert report.status_code == 200, report.text
    body = report.json()
    assert (body["pitches"], body["won"], body["lost"]) == (5, 1, 4)
    assert Decimal(body["win_rate"]) == Decimal("20.00")
    assert body["top_lost_reasons"][0] == {"reason": "fee", "count": 2}
    assert len(body["top_lost_reasons"]) == 3
    assert body["closed_deals"] == 1
    assert Decimal(body["average_duration_days"]) == Decimal("90.00")
    assert [(s["sector"], Decimal(s["average_ev_ebitda"])) for s in body["ev_ebitda_by_sector"]] == [
        ("Technology", Decimal("8.50"))
    ]

    earlier = (await get(world.partner, f"{REPORTS}/win-loss", **{"from": "2026-01-01", "to": "2026-09-30"})).json()
    assert (earlier["pitches"], earlier["closed_deals"], earlier["win_rate"]) == (0, 0, None)


# AUTHZ-02 c4


async def test_reports_and_exports_exclude_deals_i_am_not_staffed_on(world):
    hidden = await deal(world, "Kestrel", lead=world.outsider)
    await fees(world, hidden, lead=world.outsider, retainer="77777")
    await forecast(world, hidden, 100, "2026-11-15", lead=world.outsider)
    mine = await deal(world, "Falcon")
    await fees(world, mine)
    await forecast(world, mine, 100, "2026-11-15")

    forecast_body = (await get(world.partner, f"{REPORTS}/fee-forecast")).json()
    assert forecast_body["deals"] == 1
    october = forecast_body["months"][0]
    assert Decimal(october["retainers"]) == Decimal("10000.00")
    assert "77777" not in str(forecast_body)

    export = await export_(world.partner, f"{REPORTS}/fee-forecast/export")
    rows = list(load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True))
    assert rows[1][1] == 10000

    home = await get(world.partner, f"{REPORTS}/home")
    assert "Kestrel" not in home.text

    current = (await get(world.outsider, f"/api/v1/mandates/{hidden}")).json()
    closed = await post(
        world.outsider,
        f"/api/v1/mandates/{hidden}/close",
        {
            "version": current["version"],
            "outcome": "closed_won",
            "final_ev": "30000",
            "buyer": "Northwind",
            "ev_ebitda": "20",
            "closing_date": "2026-10-14",
            "tombstone_allowed": False,
        },
    )
    assert closed.status_code == 200, closed.text
    report = (await get(world.partner, f"{REPORTS}/win-loss", **{"from": "2026-10-01", "to": "2026-10-31"})).json()
    assert report["closed_deals"] == 0
    assert report["ev_ebitda_by_sector"] == []


# ROLE-01 c3


async def test_read_only_gets_aggregated_dashboards_and_nothing_else(world):
    falcon = await deal(world, "Falcon")
    await fees(world, falcon)
    contact = await post(world.partner, "/api/v1/contacts/", {"first_name": "Jan", "last_name": "Peeters"})

    dashboard = await get(world.reader, f"{REPORTS}/dashboard")
    assert dashboard.status_code == 200
    assert "weighted_fees_by_quarter" not in dashboard.json()
    assert "Falcon" not in dashboard.text

    refused = [
        "/api/v1/mandates/",
        f"/api/v1/mandates/{falcon}",
        f"/api/v1/mandates/{falcon}/forecast",
        "/api/v1/companies/",
        f"/api/v1/companies/{world.client_id}",
        "/api/v1/contacts/",
        f"/api/v1/contacts/{contact.json().get('id', 1)}",
        "/api/v1/tasks/mine",
        "/api/v1/activities/deleted",
        "/api/v1/search/?q=Acme",
        "/api/v1/origination/opportunities",
        "/api/v1/fees/",
        f"{REPORTS}/home",
        f"{REPORTS}/fee-forecast",
        f"{REPORTS}/win-loss?from=2026-01-01&to=2026-12-31",
    ]
    for path in refused:
        response = await world.reader.client.get(path, headers=world.reader.headers)
        assert response.status_code in (403, 404), f"{path}: {response.status_code}"
        assert "Falcon" not in response.text
        assert "Acme" not in response.text


async def test_win_loss_counts_other_partners_pitches_and_deals_anonymously(world):
    base = "/api/v1/origination/opportunities"

    async def lost_pitch(owner, reason: str) -> None:
        body = {
            "company_id": world.client_id,
            "type": "sell_side",
            "trigger": "succession",
            "estimated_ev": "40000",
            "probability": 30,
        }
        created = await post(owner, base, body)
        assert created.status_code == 201, created.text
        lost = await post(owner, f"{base}/{created.json()['id']}/stage", {"stage": "lost", "lost_reason": reason})
        assert lost.status_code == 200, lost.text

    async def closed_deal(code: str) -> None:
        set_clock(FrozenClock(datetime(2026, 7, 16, 9, 0, tzinfo=UTC)))
        mandate_id = await deal(world, code, lead=world.outsider)
        set_clock(FrozenClock(NOW))
        current = (await get(world.outsider, f"/api/v1/mandates/{mandate_id}")).json()
        body = {
            "version": current["version"],
            "outcome": "closed_won",
            "final_ev": "30000",
            "buyer": "Northwind",
            "ev_ebitda": "6",
            "closing_date": "2026-10-14",
            "tombstone_allowed": False,
        }
        assert (await post(world.outsider, f"/api/v1/mandates/{mandate_id}/close", body)).status_code == 200

    period = {"from": "2026-10-01", "to": "2026-10-31"}
    await lost_pitch(world.partner, "fee")
    await lost_pitch(world.outsider, "timing")
    await closed_deal("Kestrel")
    single = (await get(world.partner, f"{REPORTS}/win-loss", **period)).json()
    assert (single["pitches"], single["lost"], single["closed_deals"]) == (1, 1, 0)

    for _ in range(2):
        await lost_pitch(world.outsider, "timing")
    for code in ("Osprey", "Plover"):
        await closed_deal(code)

    response = await get(world.partner, f"{REPORTS}/win-loss", **period)
    body = response.json()
    assert (body["pitches"], body["won"], body["lost"]) == (4, 0, 4)
    assert body["top_lost_reasons"][0] == {"reason": "timing", "count": 3}
    assert body["closed_deals"] == 3
    assert [(s["deals"], Decimal(s["average_ev_ebitda"])) for s in body["ev_ebitda_by_sector"]] == [(3, Decimal("6.00"))]
    for leak in ("Kestrel", "Osprey", "Plover", "Northwind", "otto"):
        assert leak not in response.text
