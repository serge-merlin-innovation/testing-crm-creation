"""CON-02 c1, CON-03, CON-04, CON-05: search, relationship holders, merging duplicates and Excel import."""

import io
import time
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select, text

from src.infrastructure.clock import FrozenClock, SystemClock, set_clock
from src.infrastructure.jobs import job_runner
from src.modules.audit.models import AuditEntry
from src.modules.contact.models import Contact
from src.modules.data_import import service as import_service
from src.modules.data_import.jobs import IMPORT_JOB
from src.modules.data_import.service import ImportService
from src.modules.notification.models import Notification
from src.modules.notification.service import NotificationService
from src.modules.role.permissions import RoleName
from tests.integration.crm.deal_helpers import build_world, create_mandate

pytestmark = pytest.mark.asyncio

NOW = datetime(2026, 10, 10, 9, 0, tzinfo=UTC)


@pytest.fixture
async def world(client, seeded_db):
    clock = FrozenClock(NOW)
    set_clock(clock)
    built = await build_world(seeded_db)
    built.clock = clock
    built.partner = await built.actor("paula", RoleName.PARTNER)
    built.analyst = await built.actor("anna", RoleName.ANALYST)
    built.outsider = await built.actor("otto", RoleName.PARTNER)
    built.admin = await built.actor("adam", RoleName.ADMIN)
    built.compliance = await built.actor("cora", RoleName.COMPLIANCE)
    yield built
    await built.close()
    set_clock(SystemClock())


async def contact(world, actor, first, last, **extra):
    response = await actor.client.post(
        "/api/v1/contacts/",
        json={"first_name": first, "last_name": last, "source": "Forum", **extra},
        headers=actor.headers,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def search(actor, **params):
    response = await actor.client.get("/api/v1/search/", params=params, headers=actor.headers)
    assert response.status_code == 200, response.text
    return response.json()


# CON-03 c1
async def test_search_groups_by_type_and_hides_deals_i_am_not_on(world):
    client_id = await world.company("Zenith Industries")
    await contact(world, world.partner, "Zeno", "Zenithal", company_id=client_id)
    await create_mandate(world, world.partner, client_id, code_name="Zenobia")

    mine = await search(world.partner, q="zen")
    assert set(mine) == {"query", "companies", "contacts", "deals"}
    assert [c["name"] for c in mine["companies"]] == ["Zenith Industries"]
    assert [c["last_name"] for c in mine["contacts"]] == ["Zenithal"]
    assert [d["code_name"] for d in mine["deals"]] == ["Zenobia"]

    theirs = await search(world.outsider, q="zen")
    assert theirs["deals"] == []
    assert [c["name"] for c in theirs["companies"]] == ["Zenith Industries"]


async def test_search_needs_three_characters(world):
    response = await world.partner.client.get("/api/v1/search/", params={"q": "ze"}, headers=world.partner.headers)
    assert response.status_code == 422


async def test_deals_match_on_code_name_only(world):
    client_id = await world.company("Quasar Logistics")
    await create_mandate(world, world.partner, client_id, code_name="Heron")
    found = await search(world.partner, q="quasar")
    assert found["deals"] == []


# CON-03 c1 (performance)
async def test_search_answers_under_a_second_at_100k_records(world):
    await world.db.execute(
        text(
            "INSERT INTO contacts (first_name, last_name, email, language, do_not_contact, legal_basis, source, created_at) "
            "SELECT 'First' || g, 'Last' || md5(g::text), 'c' || g || '@example.com', 'en', false, "
            "'legitimate_interest', 'seed', now() FROM generate_series(1, 70000) g"
        )
    )
    await world.db.execute(
        text(
            "INSERT INTO companies (name, normalized_name, type, country_id, created_at) "
            "SELECT 'Company ' || md5(g::text), 'company ' || md5(g::text), 'target', :country, now() "
            "FROM generate_series(1, 30000) g"
        ),
        {"country": world.country_id},
    )
    await world.db.commit()
    await world.db.execute(text("ANALYZE"))
    await contact(world, world.partner, "Élodie", "Vandersmissen")
    await search(world.partner, q="warmup")
    started = time.perf_counter()
    found = await search(world.partner, q="vandersmisen")
    elapsed = time.perf_counter() - started
    assert [c["last_name"] for c in found["contacts"]][:1] == ["Vandersmissen"]
    assert elapsed < 1.0, f"search took {elapsed:.2f}s"


# CON-03 c4
@pytest.mark.parametrize("typed", ["Lefevre", "lefèbvre", "LEFEVBRE", "Lefévre"])
async def test_fuzzy_match_survives_typos_and_missing_accents(world, typed):
    await contact(world, world.partner, "Hélène", "Lefèvre")
    await world.company("Société Générale Belgique")
    found = await search(world.partner, q=typed)
    assert [c["last_name"] for c in found["contacts"]] == ["Lefèvre"]


async def test_company_found_without_accents(world):
    await world.company("Société Générale Belgique")
    found = await search(world.partner, q="societe generale")
    assert [c["name"] for c in found["companies"]] == ["Société Générale Belgique"]


# CON-03 c2
async def test_saved_filter_is_reusable_and_shareable(world):
    pe_fund = await world.company("Benelux Capital Partners", "financial_investor")
    await world.company("Rhine Strategic", "strategic_buyer")
    criteria = {"company_types": ["financial_investor"], "country_ids": [world.country_id]}
    created = await world.partner.client.post(
        "/api/v1/saved-filters/",
        json={"name": "PE funds Benelux", "criteria": criteria},
        headers=world.partner.headers,
    )
    assert created.status_code == 201, created.text
    filter_id = created.json()["id"]

    reused = await search(world.partner, filter_id=filter_id)
    assert [c["id"] for c in reused["companies"]] == [pe_fund]

    not_yet = await world.analyst.client.get("/api/v1/saved-filters/", headers=world.analyst.headers)
    assert not_yet.json() == []
    hidden = await world.analyst.client.get("/api/v1/search/", params={"filter_id": filter_id}, headers=world.analyst.headers)
    assert hidden.status_code == 404

    shared = await world.partner.client.put(
        f"/api/v1/saved-filters/{filter_id}/shared", json={"shared": True}, headers=world.partner.headers
    )
    assert shared.json()["shared"] is True
    listed = (await world.analyst.client.get("/api/v1/saved-filters/", headers=world.analyst.headers)).json()
    assert [(f["name"], f["owner_name"]) for f in listed] == [("PE funds Benelux", "paula")]
    assert [c["id"] for c in (await search(world.analyst, filter_id=filter_id))["companies"]] == [pe_fund]
    forbidden = await world.analyst.client.delete(f"/api/v1/saved-filters/{filter_id}", headers=world.analyst.headers)
    assert forbidden.status_code == 403


async def test_ticket_size_filter_narrows_deals(world):
    client_id = await world.company("Acme Holding")
    await create_mandate(world, world.partner, client_id, code_name="Small", estimated_ev="5000")
    await create_mandate(world, world.partner, client_id, code_name="Mid", estimated_ev="25000")
    found = await search(world.partner, ev_min="10000", ev_max="50000")
    assert [d["code_name"] for d in found["deals"]] == ["Mid"]


# CON-02 c1
async def test_relationship_holder_shows_on_contact_and_in_search(world):
    contact_id = await contact(world, world.partner, "Marc", "Dubois")
    response = await world.partner.client.put(
        f"/api/v1/contacts/{contact_id}/relationship-holders/me", json={"strength": 3}, headers=world.partner.headers
    )
    assert response.status_code == 200, response.text
    await world.analyst.client.put(
        f"/api/v1/contacts/{contact_id}/relationship-holders/me", json={"strength": 1}, headers=world.analyst.headers
    )
    shown = (await world.analyst.client.get(f"/api/v1/contacts/{contact_id}", headers=world.analyst.headers)).json()
    assert [(h["name"], h["strength"]) for h in shown["relationship_holders"]] == [("paula", 3), ("anna", 1)]
    found = await search(world.analyst, q="dubois")
    assert [(h["name"], h["strength"]) for h in found["contacts"][0]["relationship_holders"]] == [("paula", 3), ("anna", 1)]

    bad = await world.partner.client.put(
        f"/api/v1/contacts/{contact_id}/relationship-holders/me", json={"strength": 4}, headers=world.partner.headers
    )
    assert bad.status_code == 422


async def log_activity(actor, links):
    response = await actor.client.post(
        "/api/v1/activities/",
        json={"type": "call", "occurred_at": NOW.isoformat(), "participants": [], "note": "Call", "links": links},
        headers=actor.headers,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def duplicate_company(world, name):
    partner = world.partner
    response = await partner.client.post(
        "/api/v1/companies/",
        json={"name": name, "type": "target", "country_id": world.country_id, "ignore_similar": True},
        headers=partner.headers,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def merge(actor, **body):
    return await actor.client.post("/api/v1/merges/", json=body, headers=actor.headers)


# CON-04 c1
async def test_merging_contacts_moves_history_to_the_survivor_and_is_audited(world):
    company_id = await world.company("Acme Holding")
    survivor = await contact(world, world.partner, "Jan", "Peeters", email="jan@old.be")
    duplicate = await contact(world, world.partner, "Jan", "Peeters", email="jan@acme.be", company_id=company_id)
    shared = await log_activity(
        world.partner, [{"record_type": "contact", "record_id": survivor}, {"record_type": "contact", "record_id": duplicate}]
    )
    moved = await log_activity(world.partner, [{"record_type": "contact", "record_id": duplicate}])
    await world.partner.client.post(
        "/api/v1/notes/",
        json={"record_type": "contact", "record_id": duplicate, "body": "Prefers mornings", "visibility": "shared"},
        headers=world.partner.headers,
    )
    await world.analyst.client.put(
        f"/api/v1/contacts/{duplicate}/relationship-holders/me", json={"strength": 2}, headers=world.analyst.headers
    )

    response = await merge(
        world.partner, record_type="contact", survivor_id=survivor, duplicate_id=duplicate, fields={"email": "duplicate"}
    )
    assert response.status_code == 201, response.text

    gone = await world.partner.client.get(f"/api/v1/contacts/{duplicate}", headers=world.partner.headers)
    assert gone.status_code == 404
    kept = (await world.partner.client.get(f"/api/v1/contacts/{survivor}", headers=world.partner.headers)).json()
    assert kept["email"] == "jan@acme.be"
    assert kept["current_role"]["company_id"] == company_id
    assert [h["name"] for h in kept["relationship_holders"]] == ["anna"]
    timeline = (await world.partner.client.get(f"/api/v1/timeline/contact/{survivor}", headers=world.partner.headers)).json()
    assert sorted(a["id"] for a in timeline["activities"]) == sorted([shared, moved])
    assert [n["body"] for n in timeline["notes"]] == ["Prefers mornings"]

    entry = (await world.db.execute(select(AuditEntry).where(AuditEntry.action == "contact.merged"))).scalar_one()
    assert entry.record_id == str(survivor)
    assert entry.old_values["duplicate"]["id"] == duplicate
    assert entry.new_values["picked"] == {"email": "jan@acme.be"}


async def test_merging_companies_moves_deals_and_people(world):
    survivor = await world.company("Acme Holding")
    duplicate = await duplicate_company(world, "ACME Holdings")
    person = await contact(world, world.partner, "Jan", "Peeters", company_id=duplicate)
    deal = await create_mandate(world, world.partner, duplicate)
    response = await merge(world.partner, record_type="company", survivor_id=survivor, duplicate_id=duplicate)
    assert response.status_code == 201, response.text
    moved_deal = (await world.partner.client.get(f"/api/v1/mandates/{deal['id']}", headers=world.partner.headers)).json()
    assert moved_deal["client_id"] == survivor
    person_now = (await world.partner.client.get(f"/api/v1/contacts/{person}", headers=world.partner.headers)).json()
    assert person_now["current_role"]["company_id"] == survivor


# CON-04 c2
async def test_admin_undoes_a_merge_within_30_days(world):
    survivor = await contact(world, world.partner, "Jan", "Peeters", email="jan@old.be")
    duplicate = await contact(world, world.partner, "Jan", "Peeters", email="jan@acme.be")
    both = await log_activity(
        world.partner, [{"record_type": "contact", "record_id": survivor}, {"record_type": "contact", "record_id": duplicate}]
    )
    only_dup = await log_activity(world.partner, [{"record_type": "contact", "record_id": duplicate}])
    merged = (
        await merge(
            world.partner, record_type="contact", survivor_id=survivor, duplicate_id=duplicate, fields={"email": "duplicate"}
        )
    ).json()

    not_admin = await world.partner.client.post(f"/api/v1/merges/{merged['id']}/undo", headers=world.partner.headers)
    assert not_admin.status_code == 403

    set_clock(FrozenClock(NOW + timedelta(days=29)))
    listed = (await world.admin.client.get("/api/v1/merges/", headers=world.admin.headers)).json()
    assert [m["id"] for m in listed] == [merged["id"]]
    undone = await world.admin.client.post(f"/api/v1/merges/{merged['id']}/undo", headers=world.admin.headers)
    assert undone.status_code == 200, undone.text

    restored = (await world.partner.client.get(f"/api/v1/contacts/{duplicate}", headers=world.partner.headers)).json()
    assert restored["email"] == "jan@acme.be"
    kept = (await world.partner.client.get(f"/api/v1/contacts/{survivor}", headers=world.partner.headers)).json()
    assert kept["email"] == "jan@old.be"
    dup_timeline = (
        await world.partner.client.get(f"/api/v1/timeline/contact/{duplicate}", headers=world.partner.headers)
    ).json()
    assert sorted(a["id"] for a in dup_timeline["activities"]) == sorted([both, only_dup])
    surv_timeline = (
        await world.partner.client.get(f"/api/v1/timeline/contact/{survivor}", headers=world.partner.headers)
    ).json()
    assert [a["id"] for a in surv_timeline["activities"]] == [both]
    assert (await world.db.execute(select(AuditEntry).where(AuditEntry.action == "contact.merge_undone"))).scalar_one()

    again = await world.admin.client.post(f"/api/v1/merges/{merged['id']}/undo", headers=world.admin.headers)
    assert again.status_code == 409


async def test_undo_closes_after_30_days(world):
    survivor = await contact(world, world.partner, "Jan", "Peeters")
    duplicate = await contact(world, world.partner, "Jan", "Peeters")
    merged = (await merge(world.partner, record_type="contact", survivor_id=survivor, duplicate_id=duplicate)).json()
    set_clock(FrozenClock(NOW + timedelta(days=31)))
    late = await world.admin.client.post(f"/api/v1/merges/{merged['id']}/undo", headers=world.admin.headers)
    assert late.status_code == 409


# CON-04 c3
async def test_merge_blocked_when_a_record_is_on_a_deal_i_cannot_access(world):
    client_id = await world.company("Acme Holding")
    duplicate_id = await duplicate_company(world, "Acme Holdings")
    await create_mandate(world, world.partner, client_id)

    blocked = await merge(world.outsider, record_type="company", survivor_id=duplicate_id, duplicate_id=client_id)
    assert blocked.status_code == 403
    assert "Compliance" in blocked.text
    assert await world.db.scalar(select(func.count()).select_from(AuditEntry).where(AuditEntry.action == "company.merged")) == 0

    person = await contact(world, world.partner, "Jan", "Peeters", company_id=client_id)
    other = await contact(world, world.partner, "Jan", "Peeters")
    contact_blocked = await merge(world.outsider, record_type="contact", survivor_id=other, duplicate_id=person)
    assert contact_blocked.status_code == 403

    by_compliance = await merge(world.compliance, record_type="company", survivor_id=duplicate_id, duplicate_id=client_id)
    assert by_compliance.status_code == 201, by_compliance.text


def workbook(rows):
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def upload(actor, entity, content, name="legacy.xlsx"):
    return await actor.client.post(
        "/api/v1/imports/",
        data={"entity": entity},
        files={"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=actor.headers,
    )


CONTACT_MAPPING = {"first_name": "Voornaam", "last_name": "Achternaam", "email": "E-mail", "company": "Bedrijf"}


# CON-05 c1
async def test_mapped_import_previews_new_duplicate_and_error_rows(world):
    await world.company("Acme Holding")
    await contact(world, world.partner, "Jan", "Peeters", email="jan@acme.be")
    content = workbook(
        [
            ["Voornaam", "Achternaam", "E-mail", "Bedrijf"],
            ["Els", "Maes", "els@acme.be", "Acme Holding"],
            ["Jan", "Peeters", "JAN@acme.be", None],
            ["Piet", None, "piet@acme.be", None],
            ["Ann", "Claes", "not-an-email", None],
            ["Els", "Maes", "els@acme.be", None],
        ]
    )
    forbidden = await upload(world.partner, "contact", content)
    assert forbidden.status_code == 403
    uploaded = await upload(world.admin, "contact", content)
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["headers"] == ["Voornaam", "Achternaam", "E-mail", "Bedrijf"]
    assert uploaded.json()["row_count"] == 5

    preview = await world.admin.client.post(
        f"/api/v1/imports/{uploaded.json()['id']}/preview", json={"mapping": CONTACT_MAPPING}, headers=world.admin.headers
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["counts"] == {"new": 1, "duplicate": 2, "error": 2}
    statuses = {row["row"]: row["status"] for row in body["rows"]}
    assert statuses == {2: "new", 3: "duplicate", 4: "error", 5: "error", 6: "duplicate"}
    assert await world.db.scalar(select(func.count()).select_from(Contact)) == 1


# CON-05 c2
async def test_completed_import_offers_an_error_file(world):
    content = workbook(
        [
            ["name", "type", "country", "vat_or_registry_number"],
            ["Nordic Fund", "financial_investor", "NL", "NL001"],
            ["Bad Type", "unicorn", "NL", "NL002"],
            ["No Country", "target", "XX", "NL003"],
        ]
    )
    uploaded = (await upload(world.admin, "company", content)).json()
    mapping = {"name": "name", "type": "type", "country": "country", "vat_or_registry_number": "vat_or_registry_number"}
    await world.admin.client.post(
        f"/api/v1/imports/{uploaded['id']}/preview", json={"mapping": mapping}, headers=world.admin.headers
    )
    done = await world.admin.client.post(f"/api/v1/imports/{uploaded['id']}/confirm", headers=world.admin.headers)
    assert done.status_code == 200, done.text
    assert done.json()["status"] == "completed"
    assert done.json()["created"] == 1
    url = done.json()["error_file_url"]
    assert url == f"/api/v1/imports/{uploaded['id']}/errors.xlsx"

    download = await world.admin.client.get(url, headers=world.admin.headers)
    assert download.status_code == 200
    sheet = load_workbook(io.BytesIO(download.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Row", "name", "type", "country", "vat_or_registry_number", "Errors")
    assert [(r[0], r[1]) for r in rows[1:]] == [(3, "Bad Type"), (4, "No Country")]
    assert "type must be one of" in rows[1][5]
    assert "country 'XX'" in rows[2][5]


# CON-05 c3
async def test_large_import_runs_in_background_and_notifies(world, monkeypatch):
    monkeypatch.setattr(import_service, "BACKGROUND_THRESHOLD", 3)
    rows = [["first_name", "last_name", "email"]] + [[f"P{i}", "Bulk", f"p{i}@bulk.be"] for i in range(4)]
    uploaded = (await upload(world.admin, "contact", workbook(rows))).json()
    mapping = {"first_name": "first_name", "last_name": "last_name", "email": "email"}
    await world.admin.client.post(
        f"/api/v1/imports/{uploaded['id']}/preview", json={"mapping": mapping}, headers=world.admin.headers
    )
    queued = await world.admin.client.post(f"/api/v1/imports/{uploaded['id']}/confirm", headers=world.admin.headers)
    assert queued.json()["status"] == "queued"
    assert queued.json()["background"] is True
    assert await world.db.scalar(select(func.count()).select_from(Contact)) == 0

    assert IMPORT_JOB in job_runner.names
    await ImportService(NotificationService(world.emails)).run_queued(world.db, NOW)

    status = (await world.admin.client.get(f"/api/v1/imports/{uploaded['id']}", headers=world.admin.headers)).json()
    assert status["status"] == "completed"
    assert status["created"] == 4
    note = (await world.db.execute(select(Notification).where(Notification.user_id == world.admin.id))).scalar_one()
    assert note.kind == "data_import.completed"
    assert "4 contacts created" in note.body
    assert [e.subject for e in world.emails.sent if e.to == world.admin.email] == ["Import of legacy.xlsx finished"]


async def test_threshold_is_fifty_thousand_rows():
    assert import_service.BACKGROUND_THRESHOLD == 50_000


async def hidden_deal_with_parties(world):
    """A deal the outsider cannot see, with a buyer company, a buyer contact and a counterparty on it."""
    client_id = await world.company("Acme Holding")
    mandate = await create_mandate(world, world.partner, client_id)
    buyer_id = await world.company("Bidco Capital", "financial_investor")
    counterparty_id = await world.company("Rival Industries")
    buyer_contact = await contact(world, world.partner, "Bea", "Bidder")
    url = f"/api/v1/mandates/{mandate['id']}"
    added = await world.partner.client.post(
        f"{url}/buyers",
        json={"items": [{"company_id": buyer_id, "owner_id": world.partner.id, "rationale": "Fit"}]},
        headers=world.partner.headers,
    )
    assert added.status_code == 200, added.text
    patched = await world.partner.client.patch(
        f"{url}/buyers/{added.json()['added'][0]}", json={"contact_id": buyer_contact}, headers=world.partner.headers
    )
    assert patched.status_code == 200, patched.text
    counterparty = await world.partner.client.post(
        f"{url}/counterparties", json={"company_id": counterparty_id}, headers=world.partner.headers
    )
    assert counterparty.status_code == 201, counterparty.text
    return {"buyer": buyer_id, "contact": buyer_contact, "counterparty": counterparty_id}


# CON-04 c3: a buyer or counterparty on a hidden deal blocks the merge, and nothing moves.
@pytest.mark.parametrize("party", ["buyer", "contact", "counterparty"])
async def test_merge_blocked_for_a_party_on_a_hidden_deal(world, party):
    parties = await hidden_deal_with_parties(world)
    if party == "contact":
        other = await contact(world, world.outsider, "Bea", "Bidder")
        record_type = "contact"
    else:
        other = await duplicate_company(world, "Unrelated Duplicate")
        record_type = "company"
    blocked = await merge(world.outsider, record_type=record_type, survivor_id=other, duplicate_id=parties[party])
    assert blocked.status_code == 403, blocked.text
    assert "deal_buyers" not in blocked.text
