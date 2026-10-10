"""Ticket #13: investor profiles (CO-04), buyer long list (BUY-01/02), status grid (BUY-03), NDAs (BUY-04),
buyer status hidden from the unstaffed (AUTHZ-02 c5) and Do not contact (CON-01 c5)."""

import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import update

from src.infrastructure.clock import FrozenClock, SystemClock, set_clock
from src.modules.buyer.jobs import run_buyer_jobs
from src.modules.buyer.schemas import BuyerStatus
from src.modules.mandate.models import Mandate
from src.modules.notification.service import NotificationService
from src.modules.role.permissions import RoleName
from tests.integration.crm.deal_helpers import build_world, create_mandate

pytestmark = pytest.mark.asyncio

NOW = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)


@pytest.fixture
async def world(client, seeded_db):
    set_clock(FrozenClock(NOW))
    built = await build_world(seeded_db)
    built.partner = await built.actor("paula", RoleName.PARTNER)
    built.analyst = await built.actor("anna", RoleName.ANALYST)
    built.outsider = await built.actor("oscar", RoleName.PARTNER)
    built.client_id = await company(built, "Acme Holding", "target")
    built.target_id = await company(built, "Widget Works", "target")
    yield built
    await built.close()
    set_clock(SystemClock())


async def get(actor, path: str, **params):
    return await actor.client.get(path, headers=actor.headers, params=params)


async def send(actor, method: str, path: str, body: dict):
    return await actor.client.request(method, path, json=body, headers=actor.headers)


async def company(world, name: str, company_type: str = "financial_investor", parent_id: int | None = None) -> int:
    partner = next(iter(world.actors.values()))
    response = await send(
        partner,
        "POST",
        "/api/v1/companies/",
        {
            "name": name,
            "type": company_type,
            "country_id": world.country_id,
            "vat_or_registry_number": f"{name}-reg",
            "ignore_similar": True,
            "parent_id": parent_id,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def profile(world, company_id: int, **values):
    body = {
        "sector_ids": [world.sector_id],
        "country_ids": [world.country_id],
        "ticket_min": "10",
        "ticket_max": "50",
        "ev_ebitda_min": "5",
        "ev_ebitda_max": "9",
        "stake": "majority",
        **values,
    }
    response = await send(world.partner, "PUT", f"/api/v1/companies/{company_id}/investor-profile", body)
    assert response.status_code == 200, response.text
    return response.json()


async def deal(world) -> int:
    mandate = await create_mandate(world, world.partner, world.client_id, target_id=world.target_id)
    response = await send(
        world.partner, "POST", f"/api/v1/mandates/{mandate['id']}/team", {"user_id": world.analyst.id, "role": "analyst"}
    )
    assert response.status_code == 201, response.text
    return mandate["id"]


async def add(world, mandate_id: int, company_ids: list[int], confirm: bool = False, **extra):
    items = [{"company_id": cid, "owner_id": world.analyst.id, "rationale": "Fits the thesis", **extra} for cid in company_ids]
    return await send(world.analyst, "POST", f"/api/v1/mandates/{mandate_id}/buyers", {"items": items, "confirm": confirm})


async def buyers(world, mandate_id: int, actor=None) -> dict:
    response = await get(actor or world.analyst, f"/api/v1/mandates/{mandate_id}/buyers")
    assert response.status_code == 200, response.text
    return response.json()


async def set_status(world, mandate_id: int, buyer_ids: list[int], status: str, **extra):
    return await send(
        world.analyst,
        "POST",
        f"/api/v1/mandates/{mandate_id}/buyers/status",
        {"buyer_ids": buyer_ids, "status": status, **extra},
    )


async def approve(world, mandate_id: int, buyer_id: int, approval: str = "approved"):
    response = await send(world.partner, "PATCH", f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}", {"approval": approval})
    assert response.status_code == 200, response.text


async def listed(world, mandate_id: int, n: int = 1) -> list[int]:
    funds = [await company(world, f"Fund {i}") for i in range(n)]
    response = await add(world, mandate_id, funds)
    assert response.status_code == 200, response.text
    return [row["id"] for row in (await buyers(world, mandate_id))["buyers"]]


# CO-04


async def test_co04_c1_criteria_are_stored_and_searchable_in_suggest_buyers(world):
    fund = await company(world, "Northstar Capital")
    stored = await profile(world, fund)
    mandate_id = await deal(world)

    suggestions = (await get(world.analyst, f"/api/v1/mandates/{mandate_id}/buyers/suggestions")).json()

    assert stored["ticket_min"] == "10.00"
    assert stored["stake"] == "majority"
    assert stored["sector_ids"] == [world.sector_id]
    assert [row["company_id"] for row in suggestions] == [fund]


async def test_co04_c2_portfolio_companies_show_as_group_entities(world):
    fund = await company(world, "Northstar Capital")
    portfolio = await company(world, "Bakery Group", "target")

    linked = await send(world.partner, "POST", f"/api/v1/companies/{fund}/portfolio", {"company_id": portfolio})
    view = (await get(world.partner, f"/api/v1/companies/{fund}/investor-profile")).json()
    group = (await get(world.partner, f"/api/v1/companies/{fund}/group-members")).json()

    assert linked.status_code == 200, linked.text
    assert [entity["id"] for entity in view["portfolio"]] == [portfolio]
    assert portfolio in group


async def test_co04_c3_profile_not_updated_for_12_months_is_flagged(world):
    fund = await company(world, "Northstar Capital")
    await profile(world, fund)
    fresh = (await get(world.partner, f"/api/v1/companies/{fund}/investor-profile")).json()

    set_clock(FrozenClock(NOW + timedelta(days=366)))
    stale = (await get(world.partner, f"/api/v1/companies/{fund}/investor-profile")).json()

    assert fresh["outdated"] is False
    assert stale["outdated"] is True


async def test_co04_c4_decline_reasons_from_accessible_closed_deals_are_listed(world):
    fund = await company(world, "Northstar Capital")
    mandate_id = await deal(world)
    [buyer_id] = (await add(world, mandate_id, [fund])).json()["added"]
    assert (await set_status(world, mandate_id, [buyer_id], "declined", decline_reason="price")).status_code == 200
    version = (await get(world.partner, f"/api/v1/mandates/{mandate_id}")).json()["version"]
    closed = await send(
        world.partner,
        "POST",
        f"/api/v1/mandates/{mandate_id}/close",
        {"version": version, "outcome": "aborted", "abort_reason": "Seller withdrew"},
    )
    assert closed.status_code == 200, closed.text

    mine = (await get(world.partner, f"/api/v1/companies/{fund}/investor-profile")).json()["declines"]
    theirs = (await get(world.outsider, f"/api/v1/companies/{fund}/investor-profile")).json()["declines"]

    assert [(row["code_name"], row["reason"]) for row in mine] == [("Falcon", "price")]
    assert theirs == []


# BUY-01


async def test_buy01_c1_suggest_ranks_matching_companies_and_adds_in_bulk(world):
    best = await company(world, "Best Fit Partners")
    await profile(world, best)
    partial = await company(world, "Partial Fit Partners")
    await profile(world, partial, ticket_min="100", ticket_max="500")
    await profile(
        world, await company(world, "No Fit Partners"), sector_ids=[], country_ids=[], ticket_min="900", ticket_max="1000"
    )
    mandate_id = await deal(world)

    suggestions = (await get(world.analyst, f"/api/v1/mandates/{mandate_id}/buyers/suggestions")).json()
    ranked = [row["company_id"] for row in suggestions]
    added = await add(world, mandate_id, ranked)

    assert ranked == [best, partial]
    assert suggestions[0]["score"] > suggestions[1]["score"]
    assert added.json()["added"]
    assert len(added.json()["added"]) == 2


async def test_buy01_c2_buyer_already_on_the_list_is_skipped(world):
    fund = await company(world, "Northstar Capital")
    mandate_id = await deal(world)
    await add(world, mandate_id, [fund])

    again = await add(world, mandate_id, [fund])

    assert again.json()["added"] == []
    assert again.json()["skipped"] == [fund]
    assert len((await buyers(world, mandate_id))["buyers"]) == 1


async def test_buy01_c3_conflict_or_exclusion_list_shows_a_warning(world):
    excluded = await company(world, "Rival Industries", "strategic_buyer")
    conflicted = await company(world, "Other Side Capital")
    await create_mandate(world, world.partner, conflicted, code_name="Heron")
    mandate_id = await deal(world)
    response = await send(world.partner, "POST", f"/api/v1/companies/{world.client_id}/exclusions", {"company_id": excluded})
    assert response.status_code == 201, response.text

    warned = await add(world, mandate_id, [excluded, conflicted])
    confirmed = await add(world, mandate_id, [excluded, conflicted], confirm=True)

    reasons = {(row["company_id"], row["reason"]) for row in warned.json()["warnings"]}
    assert reasons == {(excluded, "client_exclusion"), (conflicted, "conflict")}
    assert warned.json()["added"] == []
    assert len(confirmed.json()["added"]) == 2


async def test_buy01_c4_owner_and_rationale_are_required(world):
    fund = await company(world, "Northstar Capital")
    mandate_id = await deal(world)
    url = f"/api/v1/mandates/{mandate_id}/buyers"

    no_rationale = await send(
        world.analyst, "POST", url, {"items": [{"company_id": fund, "owner_id": world.analyst.id, "rationale": " "}]}
    )
    no_owner = await send(world.analyst, "POST", url, {"items": [{"company_id": fund, "rationale": "Fit"}]})
    outsider_owner = await send(
        world.analyst, "POST", url, {"items": [{"company_id": fund, "owner_id": world.outsider.id, "rationale": "Fit"}]}
    )

    assert no_rationale.status_code == 422
    assert no_owner.status_code == 422
    assert outsider_owner.status_code == 422
    assert (await buyers(world, mandate_id))["buyers"] == []


# BUY-02


async def test_buy02_c1_excel_export_has_buyer_type_rationale_contact_without_notes(world):
    fund = await company(world, "Northstar Capital")
    contact = await send(
        world.partner,
        "POST",
        "/api/v1/contacts/",
        {"first_name": "Nina", "last_name": "North", "source": "Forum", "company_id": fund},
    )
    mandate_id = await deal(world)
    [buyer_id] = (await add(world, mandate_id, [fund], contact_id=contact.json()["id"])).json()["added"]
    await send(world.analyst, "PATCH", f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}", {"notes": "Secret view on price"})

    response = await get(world.analyst, f"/api/v1/mandates/{mandate_id}/buyers/export")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("Buyer", "Type", "Rationale", "Contact")
    assert rows[1] == ("Northstar Capital", "Financial investor / PE", "Fits the thesis", "Nina North")
    assert "Secret" not in str(rows)


async def test_buy02_c2_only_approved_buyers_can_be_contacted(world):
    mandate_id = await deal(world)
    first, second, third = await listed(world, mandate_id, 3)
    await approve(world, mandate_id, first)
    await approve(world, mandate_id, second, "excluded")

    approved = await set_status(world, mandate_id, [first], "contacted")
    excluded = await set_status(world, mandate_id, [second], "contacted")
    pending = await set_status(world, mandate_id, [third], "teaser_sent")

    assert approved.status_code == 200
    assert excluded.status_code == 422
    assert pending.status_code == 422


# BUY-03


async def test_buy03_c1_status_change_stores_date_and_user_and_funnel_updates(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await approve(world, mandate_id, buyer_id)
    before = (await buyers(world, mandate_id))["funnel"]

    set_clock(FrozenClock(NOW + timedelta(days=1)))
    await set_status(world, mandate_id, [buyer_id], "contacted")
    after = await buyers(world, mandate_id)
    row = after["buyers"][0]

    assert before["client_approved"] == 1
    assert before["contacted"] == 0
    assert after["funnel"]["client_approved"] == 0
    assert after["funnel"]["contacted"] == 1
    assert row["status_changed_at"].startswith("2026-10-10")
    assert row["status_changed_by_name"] == "anna"


async def test_buy03_c2_im_sent_requires_nda_signed(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await approve(world, mandate_id, buyer_id)

    blocked = await set_status(world, mandate_id, [buyer_id], "im_sent")
    await send(
        world.analyst,
        "PUT",
        f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}/nda",
        {"sent_date": "2026-10-01", "signed_date": "2026-10-05"},
    )
    allowed = await set_status(world, mandate_id, [buyer_id], "im_sent")

    assert blocked.status_code == 422
    assert "NDA" in blocked.text
    assert allowed.status_code == 200


async def test_buy03_c3_bulk_update_of_20_buyers_gives_all_the_same_date(world):
    mandate_id = await deal(world)
    ids = await listed(world, mandate_id, 20)
    for buyer_id in ids:
        await approve(world, mandate_id, buyer_id)

    response = await set_status(world, mandate_id, ids, "teaser_sent")
    rows = (await buyers(world, mandate_id))["buyers"]

    assert response.status_code == 200, response.text
    assert {row["status"] for row in rows} == {"teaser_sent"}
    assert len({row["status_changed_at"] for row in rows}) == 1
    assert (await buyers(world, mandate_id))["funnel"]["teaser_sent"] == 20


async def test_buy03_c4_declined_requires_a_reason_from_the_list(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)

    missing = await set_status(world, mandate_id, [buyer_id], "declined")
    unknown = await set_status(world, mandate_id, [buyer_id], "declined", decline_reason="weather")
    ok = await set_status(world, mandate_id, [buyer_id], "declined", decline_reason="timing")

    assert missing.status_code == 422
    assert unknown.status_code == 422
    assert ok.status_code == 200
    assert (await buyers(world, mandate_id))["buyers"][0]["decline_reason"] == "timing"


async def _to_marketing(world, mandate_id: int) -> None:
    await world.db.execute(update(Mandate).where(Mandate.id == mandate_id).values(stage="marketing"))
    await world.db.commit()


async def test_buy03_c5_owner_is_reminded_after_14_idle_days_in_an_active_phase(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await _to_marketing(world, mandate_id)
    notifications = NotificationService(world.emails)

    await run_buyer_jobs(world.db, NOW + timedelta(days=13), notifications)
    early = [mail for mail in world.emails.sent if "no activity" in mail.subject]
    await run_buyer_jobs(world.db, NOW + timedelta(days=15), notifications)
    await run_buyer_jobs(world.db, NOW + timedelta(days=16), notifications)

    assert early == []
    reminders = [mail for mail in world.emails.sent if "no activity" in mail.subject]
    assert [mail.to for mail in reminders] == [world.analyst.email]
    assert "Falcon" in reminders[0].subject
    assert "Acme" not in reminders[0].body


async def test_buy03_c5_no_reminder_outside_an_active_phase(world):
    mandate_id = await deal(world)
    await listed(world, mandate_id)

    await run_buyer_jobs(world.db, NOW + timedelta(days=30), NotificationService(world.emails))

    assert [mail for mail in world.emails.sent if "no activity" in mail.subject] == []


# BUY-04


async def test_buy04_c1_nda_dates_and_link_move_to_nda_signed_only_on_signed_date(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await approve(world, mandate_id, buyer_id)
    url = f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}/nda"

    pending = await send(
        world.analyst,
        "PUT",
        url,
        {
            "sent_date": "2026-10-01",
            "markup_received_date": "2026-10-03",
            "agreed_date": "2026-10-04",
            "document_url": "https://docs.example.com/nda",
        },
    )
    status_before = (await buyers(world, mandate_id))["buyers"][0]["status"]
    manual = await set_status(world, mandate_id, [buyer_id], "nda_signed")
    signed = await send(world.analyst, "PUT", url, {**pending.json(), "signed_date": "2026-10-06"})
    row = (await buyers(world, mandate_id))["buyers"][0]

    assert pending.status_code == 200, pending.text
    assert pending.json()["document_url"] == "https://docs.example.com/nda"
    assert status_before == "nda_sent"
    assert manual.status_code == 422
    assert signed.status_code == 200, signed.text
    assert row["status"] == "nda_signed"
    assert row["nda"]["signed_date"] == "2026-10-06"


async def test_buy04_c2_nda_pending_more_than_10_days_reminds_the_owner(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await approve(world, mandate_id, buyer_id)
    await send(world.analyst, "PUT", f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}/nda", {"sent_date": "2026-10-09"})
    notifications = NotificationService(world.emails)

    await run_buyer_jobs(world.db, NOW + timedelta(days=10), notifications)
    early = [mail for mail in world.emails.sent if "NDA" in mail.subject]
    await run_buyer_jobs(world.db, NOW + timedelta(days=11), notifications)
    await run_buyer_jobs(world.db, NOW + timedelta(days=12), notifications)

    assert early == []
    reminders = [mail for mail in world.emails.sent if "NDA" in mail.subject]
    assert [mail.to for mail in reminders] == [world.analyst.email]


async def test_buy04_c3_standstill_and_non_solicit_end_dates_show_on_the_buyer_company(world):
    fund = await company(world, "Northstar Capital")
    mandate_id = await deal(world)
    [buyer_id] = (await add(world, mandate_id, [fund])).json()["added"]
    await approve(world, mandate_id, buyer_id)

    response = await send(
        world.analyst,
        "PUT",
        f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}/nda",
        {
            "sent_date": "2026-10-01",
            "signed_date": "2026-10-05",
            "standstill_end": "2027-10-05",
            "non_solicit_end": "2028-04-05",
        },
    )
    view = (await get(world.outsider, f"/api/v1/companies/{fund}/investor-profile")).json()

    assert response.status_code == 200, response.text
    assert view["standstill_end"] == "2027-10-05"
    assert view["non_solicit_end"] == "2028-04-05"
    assert "Falcon" not in str(view)


# AUTHZ-02 c5


async def test_authz02_c5_unstaffed_user_does_not_see_buyer_status_on_a_contact(world):
    fund = await company(world, "Northstar Capital")
    contact = (
        await send(
            world.partner,
            "POST",
            "/api/v1/contacts/",
            {"first_name": "Nina", "last_name": "North", "source": "Forum", "company_id": fund},
        )
    ).json()
    mandate_id = await deal(world)
    await add(world, mandate_id, [fund], contact_id=contact["id"])

    staffed = await get(world.analyst, f"/api/v1/contacts/{contact['id']}/buyer-roles")
    unstaffed = await get(world.outsider, f"/api/v1/contacts/{contact['id']}/buyer-roles")
    contact_view = await get(world.outsider, f"/api/v1/contacts/{contact['id']}")

    assert [row["status"] for row in staffed.json()] == ["long_list"]
    assert unstaffed.status_code == 200
    assert unstaffed.json() == []
    assert "long_list" not in contact_view.text
    assert "Falcon" not in contact_view.text


# CON-01 c5


async def test_con01_c5_do_not_contact_blocks_buyer_outreach(world):
    fund = await company(world, "Northstar Capital")
    blocked = (
        await send(
            world.partner,
            "POST",
            "/api/v1/contacts/",
            {"first_name": "Dan", "last_name": "Nope", "source": "Forum", "company_id": fund, "do_not_contact": True},
        )
    ).json()
    mandate_id = await deal(world)

    added = await add(world, mandate_id, [fund], confirm=True, contact_id=blocked["id"])

    assert added.status_code == 422
    assert "Do not contact" in added.text
    assert (await buyers(world, mandate_id))["buyers"] == []


# Merge (CON-04) carries the buyer slice along.


async def test_merging_a_buyer_company_moves_its_profile_and_deal_buyer_rows(world):
    survivor = await company(world, "Northstar Capital")
    duplicate = await company(world, "Northstar Capital Partners")
    await profile(world, duplicate)
    mandate_id = await deal(world)
    await add(world, mandate_id, [duplicate])

    response = await send(
        world.partner,
        "POST",
        "/api/v1/merges/",
        {"record_type": "company", "survivor_id": survivor, "duplicate_id": duplicate},
    )
    view = (await get(world.partner, f"/api/v1/companies/{survivor}/investor-profile")).json()

    assert response.status_code == 201, response.text
    assert [row["company_id"] for row in (await buyers(world, mandate_id))["buyers"]] == [survivor]
    assert view["has_profile"] is True
    assert view["sector_ids"] == [world.sector_id]


# Epic 7 header


SPEC_STATUSES = [
    "long_list",
    "client_approved",
    "contacted",
    "teaser_sent",
    "nda_sent",
    "nda_signed",
    "im_sent",
    "indicative_offer",
    "data_room_access",
    "binding_offer",
    "preferred_bidder",
    "declined",
]


async def test_epic7_buyer_statuses_are_the_twelve_from_the_spec(world):
    mandate_id = await deal(world)
    await listed(world, mandate_id)

    assert [status.value for status in BuyerStatus] == SPEC_STATUSES
    assert list((await buyers(world, mandate_id))["funnel"]) == SPEC_STATUSES


async def test_epic7_client_approval_moves_the_buyer_to_client_approved(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)

    await approve(world, mandate_id, buyer_id)

    assert (await buyers(world, mandate_id))["buyers"][0]["status"] == "client_approved"


async def test_epic7_data_room_access_and_preferred_bidder_need_a_signed_nda(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    await approve(world, mandate_id, buyer_id)

    blocked = await set_status(world, mandate_id, [buyer_id], "data_room_access")
    await send(
        world.analyst,
        "PUT",
        f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}/nda",
        {"sent_date": "2026-10-01", "signed_date": "2026-10-05"},
    )
    room = await set_status(world, mandate_id, [buyer_id], "data_room_access")
    preferred = await set_status(world, mandate_id, [buyer_id], "preferred_bidder")

    assert blocked.status_code == 422
    assert room.status_code == 200, room.text
    assert preferred.status_code == 200, preferred.text
    assert (await buyers(world, mandate_id))["funnel"]["preferred_bidder"] == 1


async def test_buy03_c5_an_activity_logged_on_the_buyer_resets_the_idle_clock(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    buyer = (await buyers(world, mandate_id))["buyers"][0]
    await _to_marketing(world, mandate_id)
    notifications = NotificationService(world.emails)

    set_clock(FrozenClock(NOW + timedelta(days=10)))
    logged = await send(
        world.analyst,
        "POST",
        "/api/v1/activities/",
        {
            "type": "call",
            "occurred_at": (NOW + timedelta(days=10)).isoformat(),
            "note": "Buyer call",
            "links": [{"record_type": "company", "record_id": buyer["company_id"]}],
        },
    )
    assert logged.status_code == 201, logged.text
    await run_buyer_jobs(world.db, NOW + timedelta(days=15), notifications)
    quiet = [mail for mail in world.emails.sent if "no activity" in mail.subject]
    await run_buyer_jobs(world.db, NOW + timedelta(days=25), notifications)

    assert quiet == []
    assert len([mail for mail in world.emails.sent if "no activity" in mail.subject]) == 1


async def test_co04_c1_ev_ebitda_and_stake_filter_suggest_buyers(world):
    majority = await company(world, "Majority Partners")
    await profile(world, majority)
    minority = await company(world, "Minority Partners")
    await profile(world, minority, stake="minority", ev_ebitda_min="10", ev_ebitda_max="14")
    either = await company(world, "Flexible Partners")
    await profile(world, either, stake="both", ev_ebitda_min="4", ev_ebitda_max="12")
    mandate_id = await deal(world)
    path = f"/api/v1/mandates/{mandate_id}/buyers/suggestions"

    def ids(response):
        assert response.status_code == 200, response.text
        return sorted(row["company_id"] for row in response.json())

    assert ids(await get(world.analyst, path)) == sorted([majority, minority, either])
    assert ids(await get(world.analyst, path, stake="majority")) == sorted([majority, either])
    assert ids(await get(world.analyst, path, stake="minority")) == sorted([minority, either])
    assert ids(await get(world.analyst, path, ev_ebitda="11")) == sorted([minority, either])
    assert ids(await get(world.analyst, path, ev_ebitda="6", stake="majority")) == sorted([majority, either])
    assert (await get(world.analyst, path, stake="both")).status_code == 422


async def test_con01_c5_do_not_contact_contact_cannot_be_set_on_a_listed_buyer(world):
    mandate_id = await deal(world)
    [buyer_id] = await listed(world, mandate_id)
    response = await send(
        world.partner,
        "POST",
        "/api/v1/contacts/",
        {"first_name": "Vera", "last_name": "Stil", "source": "Forum", "do_not_contact": True},
    )
    assert response.status_code == 201, response.text

    patched = await send(
        world.analyst, "PATCH", f"/api/v1/mandates/{mandate_id}/buyers/{buyer_id}", {"contact_id": response.json()["id"]}
    )

    assert patched.status_code == 422
    assert "Do not contact" in patched.text
    assert (await buyers(world, mandate_id))["buyers"][0]["contact_id"] is None
